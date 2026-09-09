from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).with_name("ALIS人机民族志重建_审阅版.docx")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = "000000", size: int = 9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color: str = "B7C3D0", size: str = "4") -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def style_run(run, *, bold: bool = False, italic: bool = False, color: str = "000000", size: int = 10) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_para(doc: Document, text: str = "", *, style: str | None = None, align=None, first_line: bool = True, space_after: int = 5) -> None:
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.line_spacing = 1.28
    p.paragraph_format.space_after = Pt(space_after)
    if first_line and text and style is None:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    style_run(run, size=10)


def add_bullet(doc: Document, text: str, *, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Cm(0.75 + 0.45 * level)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.line_spacing = 1.22
    p.paragraph_format.space_after = Pt(3)
    style_run(p.add_run(text), size=9.5)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    style_run(run, bold=True, color="111111", size=15 if level == 1 else 11.5)


def add_code(doc: Document, lines: str) -> None:
    for line in lines.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.right_indent = Cm(0.35)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(line)
        style_run(run, color="31465A", size=8.5)


def add_table(doc: Document, headers: Sequence[str], rows: Iterable[Sequence[str]], widths: Sequence[float] | None = None, font_size: int = 8) -> None:
    data = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        set_cell_text(header.cells[i], value, bold=True, color="FFFFFF", size=font_size)
        set_cell_shading(header.cells[i], "355C7D")
    for row_values in data:
        row = table.add_row()
        for i, value in enumerate(row_values):
            set_cell_text(row.cells[i], str(value), size=font_size)
            if len(table.rows) % 2 == 0:
                set_cell_shading(row.cells[i], "F2F6F8")
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(doc: Document, label: str, text: str, fill: str = "EAF2F8", *, spacer: bool = True) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="8AA9BF", size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(label + "：")
    style_run(r1, bold=True, color="24445B", size=9.5)
    r2 = p.add_run(text)
    style_run(r2, size=9.5)
    if spacer:
        doc.add_paragraph().paragraph_format.space_after = Pt(0)


def setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10)
    for name, size, color in (("Heading 1", 15, "111111"), ("Heading 2", 11.5, "111111"), ("Heading 3", 10.5, "355C7D")):
        s = styles[name]
        s.font.name = "Microsoft YaHei"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)


def configure_section(section) -> None:
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.85)
    section.right_margin = Cm(1.85)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.75)


def title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(62)
    r = p.add_run("ALIS")
    style_run(r, bold=True, color="234E70", size=34)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(18)
    r = p2.add_run("AI 素养调查项目\n人机民族志重建与审阅稿")
    style_run(r, bold=True, color="111111", size=20)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(36)
    r = p3.add_run("从现存代码、阶段材料、反馈文本与协作档案恢复\n项目迭代、prompt/loop、模型角色与作者思想对齐")
    style_run(r, color="536878", size=11)
    add_callout(doc, "审阅边界", "本文把可见事实、历史记录、分析性推断和未知项分开。推断用于恢复可能的协作机制，不冒充缺失的原始对话。", "EDF5F8")
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(72)
    r = p4.add_run("项目：ALIS · AI Literacy Interview Survey\n版本：审阅版 / 2026-09-09\n作者：赖培胜与 AI 协作产出之证据重建")
    style_run(r, color="536878", size=10)
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    add_heading(doc, "阅读说明", 1)
    add_para(doc, "这不是一份替作者补写“唯一真实历史”的叙事，而是一份可回到项目文件夹逐项核对的重建稿。由于 Claude Code 的历史交互在当前环境中不能直接从原平台恢复，本文优先使用仍保存的 ALIS 原始交互 DOCX、S0-S6 文档、当前源代码、测试结果和反馈文本。涉及 Claude Code 与 DeepSeek 的一般生成特点时，另以官方公开资料作背景，不把公开能力描述改写成项目内的实际日志。")
    add_table(
        doc,
        ["证据标记", "本文如何使用"],
        [
            ("已验证事实", "当前文件、当前代码、当前测试或可直接读取的阶段报告支持；可复查。"),
            ("历史记录", "原始协作档案或阶段材料当时写下的内容；说明时间点，不保证代表今天。"),
            ("分析性推断", "把多个事实串成可能的机制；必须保留‘可能/倾向/或许’语气。"),
            ("未知/待核验", "材料不足；不填空，不把顺利的叙事当作证据。"),
        ],
        widths=[3.0, 12.1],
    )
    add_heading(doc, "目录", 1)
    for item in [
        "一、恢复结论：ALIS 当时在做什么",
        "二、迭代状况与产品分叉",
        "三、prompt 与 loop：可见设计和分析性重建",
        "四、Claude Code：工具化结对开发的生成特点",
        "五、DeepSeek：语言渲染层、实际证据与公开特征",
        "六、人机民族志：角色、权力与反馈循环",
        "七、与《The Final》的思想对齐",
        "八、当前技术地图、风险和作者审阅问题",
        "附录：证据索引与重建模板",
    ]:
        add_bullet(doc, item)


def section_one(doc: Document) -> None:
    add_heading(doc, "一、恢复结论：ALIS 当时在做什么", 1)
    add_para(doc, "ALIS 不是单纯把一个问卷做成网页。现存总纲把它写成社会学研究软件：量表用于建立可追踪的测量骨架，结果页用于提供暂时的解释入口，追问和访谈用于让被试的经验、反思和修正成为数据，促进者则保留现场判断。核心铁律是“LLM 只负责语言，不负责逻辑”：评分、分支、目标选择和研究编码由确定性代码掌管。")
    add_callout(doc, "核心判断", "当前项目最稳固的连续主线不是某个 UI 版本，而是‘测量 → 解释 → 追问 → 叙事 → 研究导出’这条可回溯链路。", "EAF3EA")
    add_heading(doc, "1.1 叠加的两个产品", 2)
    add_table(
        doc,
        ["产品层", "目的", "当前代码证据", "民族志含义"],
        [
            ("评估/画像", "把 15 个陈述转成维度、轴或类型码", "`pack.ts`、`scoring.ts`、`norms.ts`、`typecode.ts`、`styleClassifier.ts`", "把模糊的 AI 素养经验先压缩为可谈论的对象。"),
            ("Interview OS 片段", "围绕结果进入六阶段访谈与反思", "`App.tsx`、`InterviewScreen.tsx`、`interview/save`", "把问卷后的解释权重新交还给被试、促进者和现场材料。"),
        ],
        widths=[2.1, 4.0, 5.0, 4.0],
        font_size=7,
    )
    add_heading(doc, "1.2 不能被抹平的双重状态", 2)
    add_para(doc, "历史 S0-S6 报告更接近详细评分内核的工程推进，而当前 `App.tsx` 的默认入口已经偏向简单模式。简单模式是 7 点量表、三组风格轴、8 类分类和六阶段访谈；详细模式是 6 点量表、7 维连续剖面、四轴类型码和一个偏差题追问。它们共用 YAML 和后端能力，却没有完全统一的转换。把二者叙述为“同一个评分算法的两个皮肤”会丢失重要事实。")
    add_para(doc, "这也解释了一个月后的记忆困难：项目不是线性地从 A 版本升级到 B 版本，而是在研究目标、产品体验、甲方反馈和工程约束之间形成了分叉。人机民族志的任务不是替这种分叉制造整齐的成功故事，而是保留它如何出现、谁在何时推动它、以及哪些决定尚未真正闭合。")


def section_two(doc: Document) -> None:
    add_heading(doc, "二、迭代状况与产品分叉", 1)
    add_heading(doc, "2.1 S0-S6：从思想铁律到可采集基线", 2)
    add_table(
        doc,
        ["阶段", "可见目标", "历史快照", "对后续的影响"],
        [
            ("S0", "ADR、schema、量表边界和铁律", "用户接受；内容可到七月底讨论", "确定 LLM/状态机分工与研究边界。"),
            ("S1", "数据地基", "13 backend tests；7 张初始表", "统一时间线和 session bundle 成为后续承载。"),
            ("S2", "评分内核", "24 frontend tests；详细 FSM", "连续剖面、区间和类型轴进入实现。"),
            ("S3", "追问闭环", "31 backend tests；fallback", "偏差题选择与一句自然追问分离。"),
            ("S4", "采集端", "40 backend / 36 frontend", "登录、PWA、录入与真实模型调用被放入体验。"),
            ("S5", "促进者与汇聚", "51 backend / 36 frontend", "保存、反应、导出、研究汇聚形成闭环。"),
            ("S6", "硬化与内测", "59 backend / 36 frontend；并发演练", "开始把‘能跑’转成可恢复、可导出的工程基线。"),
        ],
        widths=[1.4, 4.0, 4.4, 5.3],
        font_size=7,
    )
    add_para(doc, "这些历史数字与当前测试增长并不矛盾：它们属于不同阶段的快照。真正需要谨慎的是报告中“实机验证通过”的范围——它可以支持当时的阶段交付叙述，却不能自动证明今天的线上环境、密钥、浏览器或部署仍然有效。")
    add_heading(doc, "2.2 反馈如何改变问题定义", 2)
    add_para(doc, "反馈目录的共同趋势是把问题从“某个按钮坏了”推回“被试到底应该看到怎样的认知空间”。例如，早期关注计算逻辑与题目反馈，随后出现四阶段叙事、促进者引导、思考气泡、四窗口 schema，再后来转向开放空间、认知地图、节奏和目标对齐。这里的变化不是单纯视觉升级，而是对“AI 素养测量是否应该让人重新理解自己”的重新定义。")
    add_callout(doc, "反馈治理经验", "反馈材料明确警告：AI 可能把问题描述当成 bug list，随后在既有补丁上继续加补丁。更稳妥的 loop 是：接收反馈 → 重述目标 → 提出 patch/重写选择 → 作者批准 → 编辑 → 定向验证。", "FFF5E6")
    add_heading(doc, "2.3 当前分叉的审阅表", 2)
    add_table(
        doc,
        ["问题", "当前事实", "不能直接推断的内容", "需要作者决定"],
        [
            ("默认入口", "`AppMode` 默认 `simple`", "详细内核不是用户默认体验", "默认要不要改回统一内核？"),
            ("量表", "YAML 同时含 1–6 `likert` 与 1–7 `likert_simple`", "两套刻度可以无损互换", "学术报告采用哪一套作为主测量？"),
            ("画像", "简单模式 3 轴/8 类型；详细模式 7 维/4 轴", "两种类型码具有同一解释效度", "保留双层还是建立显式映射？"),
            ("追问", "详细路径有 `ProbeTarget`；简单路径进入六阶段访谈", "二者的第一问就是同一题", "访谈入口是否需要统一证据链？"),
        ],
        widths=[2.4, 4.5, 5.0, 3.2],
        font_size=7,
    )


def section_three(doc: Document) -> None:
    add_heading(doc, "三、prompt 与 loop：可见设计和分析性重建", 1)
    add_heading(doc, "3.1 可见的 prompt 设计", 2)
    add_para(doc, "后端 `probe.py` 给出了相对清晰的 prompt contract。系统侧注入角色、铁律、设计哲学、理论基础、陈述类型和 probe 风格；用户侧注入题干、维度、量表方向、类型信息和 skeleton。要求是生成一个自然、具体、与目标一致的问题，而不是重新评分、改变分支或暴露内部实现。输出超过 300 字、为空或调用失败时进入 fallback。")
    add_table(
        doc,
        ["层", "输入/约束", "谁负责", "可审计结果"],
        [
            ("内容层", "题干、维度、题型、追问骨架", "YAML 与后端 pack 校验", "版本化内容包"),
            ("逻辑层", "raw/effective、偏差策略、锁定状态", "前端纯函数 + XState + 后端校验", "确定性分数、目标、状态"),
            ("语言层", "结构化 ProbeTarget 与提示模板", "LiteLLM/DeepSeek 或 Stub", "文本、fallback、事件记录"),
            ("研究层", "回答、反应、叙事、时间线、导出", "FastAPI/SQLAlchemy/导出管线", "session bundle、CSV、manifest/hash"),
        ],
        widths=[2.2, 5.3, 4.0, 3.6],
        font_size=7,
    )
    add_heading(doc, "3.2 迭代 loop 的可能形态", 2)
    add_code(doc, "反馈/目标\n   ↓\n人：确定研究问题与验收边界\n   ↓\nAI：只读盘点代码、内容和证据 → 提出影响面与最小方案\n   ↓\n人：选择局部 patch / 定向重写 / 暂不改\n   ↓\nAI：小步编辑 → 单元测试/类型检查/构建 → 回报差异\n   ↓\n人：技术 + 学术双重审阅 → 接受、回滚或形成下一轮反馈")
    add_para(doc, "上图是“分析性重建”，不是从缺失日志逐字还原的原始 prompt。它由 S0-S6 的双重审阅制度、反馈文件的先重述后修改原则、当前代码的依赖注入/纯函数/测试结构共同支持。其价值在于说明 ALIS 的开发不是单向生成，而是人把研究边界和价值判断放回 loop 中，AI 把这些边界转成可执行代码与报告。")
    doc.add_page_break()
    add_heading(doc, "3.3 一个可复查的伪 prompt", 2)
    add_callout(doc, "证据等级：分析性重建", "以下不是原始对话引用，只是把当前 `probe.py` 的结构化输入翻译成便于审阅的模板。真实交互应以本地原始档案为准。", "FFF5E6")
    add_code(doc, "SYSTEM\n  你是 ALIS 访谈语言层。遵守：LLM 只生成语言；不重算分数；不改变目标维度；只输出一句自然追问。\n  背景：题目类型、研究理论、probe 风格、伦理边界。\n\nUSER\n  题目：{item_stem}\n  维度：{dimension_id}\n  方向：{raw/effective direction}\n  目标：{ProbeTarget}\n  骨架：{skeleton}\n  请生成一句具体、非评判、可继续回答的问题。\n\nPOST-CHECK\n  空值/过长/异常 → skeleton fallback；target 永不由文本改写。")
    add_heading(doc, "3.4 loop 中的两个张力", 2)
    add_bullet(doc, "速度与沉淀：Claude Code 可以快速扫目录、跑测试、生成补丁；如果每轮只追求“红灯变绿”，研究定义容易被局部实现带走。")
    add_bullet(doc, "一致性与真实目标：模型倾向维护当前代码的自洽性，因此会继续兼容既有命名和补丁；而反馈可能是在说“目标已变”，不是说“这里再加一个 if”。")
    add_bullet(doc, "确定性与开放性：ALIS 把评分/分支固定，把访谈叙事开放出来。这样既保留可检验性，也承认人的经验不能被一个类型码耗尽。")


def section_four(doc: Document) -> None:
    add_heading(doc, "四、Claude Code：工具化结对开发的生成特点", 1)
    add_para(doc, "项目原始档案的转录标签中，Claude 回复、思考、工具调用和工具结果数量都很大；这说明当时的协作不是一个问答窗口里的连续写作，而是带终端、文件读取、测试、构建和上下文压缩的长程代理工作。下面的“特点”分为项目内证据和公开资料背景。")
    add_table(
        doc,
        ["特点", "ALIS 内可见证据", "可能带来的好处", "必须警惕"],
        [
            ("工具先行", "大量工具调用/结果；先读规划、代码和反馈再交付", "可把抽象要求落到文件、测试和路径", "工具输出截断或压缩后，局部事实可能失真"),
            ("长程状态维护", "S0-S6、阶段 gate、继续 S3/S5/S6 的连续协作", "能把架构意图跨多轮保持", "上下文重复与 compaction 会制造“已完成”错觉"),
            ("自洽式补丁", "反馈文档明确记录 onion-style patches 风险", "对既有接口兼容，短期降低破坏面", "可能修症状而非重新对齐目标"),
            ("即时自修复", "测试失败、PowerShell/路径/依赖问题被继续处理", "把开发环境摩擦纳入 loop", "修复命令本身可能扩大范围，需授权与边界"),
            ("双重解释", "技术交付报告同时要求学术/技术复盘", "代码与研究意义不完全脱节", "流畅解释不等于学术证据"),
        ],
        widths=[2.3, 5.0, 4.0, 4.1],
        font_size=7,
    )
    add_heading(doc, "4.1 典型案例：失败如何成为协作材料", 2)
    add_para(doc, "现存历史材料记录过几类环境错误：SSH 命令中的反斜杠问题、Python 找不到 `alis` 模块、scp 路径问题，以及后期部署/隧道操作的环境差异。它们并非 ALIS 的研究结论，却很适合做人机民族志案例：AI 会把错误翻译成下一条命令，人的授权边界决定这条命令能否继续；当命令涉及远程服务器、认证、数据库或线上切换时，技术“能做”不等于“应该做”。")
    add_heading(doc, "4.2 典型案例：从 bug list 拉回目标", 2)
    add_para(doc, "727反馈1.0 直接指出，过去的 AI 可能把“问题描述 + 代码请求”当成 bug list，进而补症状。文件提出新的顺序：接收反馈、重新陈述目标、讨论 patch 还是定向重写、获得确认后再改。这个案例说明作者并没有把 Claude Code 当作拥有最终产品判断的工程师，而是在逐步建构一种“AI 生成、作者定标”的治理形式。")
    add_heading(doc, "4.3 与 Claude Code 公开使用方式的对照", 2)
    add_para(doc, "Anthropic 的公开资料把 Claude Code 描述为能直接使用终端、文件与开发工具的代码代理，并提供继续/恢复、权限模式和结构化输出等 CLI 机制。它也强调长上下文代理要保存状态、控制上下文和避免不必要的过度工程。这个背景与 ALIS 档案中的工具密集、长程、会压缩的协作形态相吻合，但它只能说明工具的设计倾向，不能替代 ALIS 本身的原始记录。")
    add_para(doc, "对作者而言，最有价值的不是“Claude 写得快”，而是把每轮写作变成了可观察事件：读了哪些文件、跑了什么测试、改了什么边界、怎样解释失败。这些事件应与人的确认、拒绝、沉默和重新定标一起归档。")


def section_five(doc: Document) -> None:
    add_heading(doc, "五、DeepSeek：语言渲染层、实际证据与公开特征", 1)
    add_para(doc, "ALIS 后端通过 LiteLLM 抽象 `LLMClient`，默认模型配置为 `deepseek/deepseek-chat`；当前实现对输出做空值/长度检查，失败回退 skeleton。也就是说，DeepSeek 在设计上是追问的语言层，不是评分或研究分类层。")
    add_callout(doc, "证据边界", "阶段 S4 报告历史上声称做过真实 DeepSeek E2E；当前审计没有重新调用外部模型，也无法从代码单独证明当时每次输出的内容、温度、延迟或成本。以下对模型特征的描述是公开资料背景，不是 ALIS 的实测报告。", "FFF5E6")
    add_heading(doc, "5.1 项目内可确认的角色", 2)
    add_table(
        doc,
        ["环节", "DeepSeek 被允许做什么", "被明确禁止/隔离的什么"],
        [
            ("追问生成", "把结构化目标和 skeleton 变成一句自然问题", "改变 item、dimension、分支或评分"),
            ("矛盾追问", "把被试文本转成可继续谈的语言", "把文本当作新的事实分数或诊断结论"),
            ("异常处理", "在正常调用时提供候选语言", "异常时必须接受 fallback，不得阻塞采集逻辑"),
            ("研究输出", "间接改善可读性", "直接决定研究编码、类型或伦理判断"),
        ],
        widths=[3.1, 6.0, 6.4],
        font_size=8,
    )
    add_heading(doc, "5.2 公开材料提供的解释框架", 2)
    add_para(doc, "DeepSeek 官方公开资料介绍了其模型架构、训练规模和推理/生成接口；API 文档也公开了温度、最大 token、JSON 输出等参数。对 ALIS 的直接启发是：模型输出必须被当成受约束的外部语言结果，而不是可信的逻辑引擎；如果未来要求结构化返回，应明确 JSON contract、空输出和 schema 失败的 fallback，并保留原始响应的最小必要审计信息。当前 `probe.py` 仍主要走文本 + 长度守卫，并没有在代码中把 JSON mode 作为追问 contract。")
    add_heading(doc, "5.3 Claude Code 与 DeepSeek 的分工对照", 2)
    add_table(
        doc,
        ["主体", "在 ALIS 里的可能位置", "生成优势", "作者需要守住的边界"],
        [
            ("Claude Code", "开发 loop 的代理/结对程序员", "读文件、写代码、跑测试、解释错误、跨阶段整理", "权限、范围、历史真实性与研究判断不能外包"),
            ("DeepSeek", "运行时追问的语言渲染器", "以较低耦合把结构化目标变成自然语言", "不让文本反向改逻辑；保留 fallback 与可追踪 target"),
            ("作者/促进者", "研究主语、验收者、现场见证者", "决定问题是否值得问、解释是否可接受、何时停止", "不能把流畅度当成真理，也不能让类型码吞掉经验"),
        ],
        widths=[3.0, 4.7, 4.5, 4.0],
        font_size=7,
    )
    add_para(doc, "由此可见，两个 AI 并不是“一个负责前端、一个负责后端”的简单分工：Claude Code 参与的是开发过程的组织与物化，DeepSeek 参与的是被试运行时的一小段语言事件。二者都被放在人的研究目的之下，但风险形式不同：前者是权限与补丁累积，后者是语言流畅性掩盖逻辑越界。")


def section_six(doc: Document) -> None:
    add_heading(doc, "六、人机民族志：角色、权力与反馈循环", 1)
    add_heading(doc, "6.1 不是“人写、AI执行”的单线叙事", 2)
    add_para(doc, "ALIS 的档案更像一个协作现场：作者提出研究问题、边界和审阅要求；Claude Code 把目录、代码、测试和部署动作组织起来；DeepSeek 在运行时生成面向被试的句子；甲方/外部测试者把体验中的断点反馈回来；作者再把反馈翻译成学术目标或产品决策。每一次翻译都会改变问题的形状。")
    add_code(doc, "研究意图 ──(作者定标)──> 工程约束 ──(Claude Code)──> 可运行结构\n    ▲                                      │                         │\n    │                                      │ 测试/错误/补丁             │ DeepSeek 语言事件\n    │                                      ▼                         ▼\n外部反馈/被试体验 <────(促进者/研究者解释)──── 访谈记录/反思 <──── 被试")
    add_para(doc, "这里存在一种不对称权力：模型可以快速生成大量可执行文本和补丁，但不能承担研究责任；作者拥有最终边界，却也可能被模型的速度、自洽和完整语气牵引。因而“人机民族志”不只记录模型输出，还记录谁能发起改变、谁能拒绝、哪个建议被沉默、哪些错误被转成下一轮规则。")
    add_heading(doc, "6.2 三类应保留的协作事件", 2)
    add_bullet(doc, "定标事件：作者把“AI 素养”“反思”“Interview OS”等抽象概念压缩成题目、维度、状态和验收条件。")
    add_bullet(doc, "越界事件：模型或部署动作接近真实密钥、线上服务器、研究编码或用户隐私时，权限与停线成为设计的一部分。")
    add_bullet(doc, "回收事件：反馈把某个局部 bug 重新解释为产品目标问题，作者决定停止加补丁、改写边界或保留未完成状态。")
    add_heading(doc, "6.3 可能的协作节奏", 2)
    add_table(
        doc,
        ["时刻", "人做什么", "AI 可能做什么", "需要留下的证据"],
        [
            ("开始", "给出研究目的和限制", "扫描目录、找入口和依赖", "范围、只读边界、证据清单"),
            ("建模", "决定概念是否可测", "提出 schema、prompt、状态图", "决策理由与未决分叉"),
            ("实现", "批准最小变更", "编辑、测试、修复环境问题", "diff、命令、测试结果"),
            ("运行", "作为促进者观察体验", "生成追问或 fallback", "target、输出、反应、错误"),
            ("复盘", "解释经验与责任", "归纳模式、提出下一轮", "事实/推断/未知分栏"),
        ],
        widths=[2.0, 4.2, 5.0, 5.0],
        font_size=7,
    )


def section_seven(doc: Document) -> None:
    add_heading(doc, "七、与《The Final》的思想对齐", 1)
    add_para(doc, "这一节不是把《The Final》当作 ALIS 的宣传语，而是用作者此前形成的思想/文风约束检查：ALIS 的工程化是否仍然保留作者的主体性、跨学科跃迁和未完成性。对齐来自此前合作中已明确的原则，因此仍应由作者逐句审阅。")
    add_heading(doc, "7.1 一条可工作的思想链", 2)
    add_code(doc, "具体经验/技术痕迹\n        ↓\n物质—社会—历史条件\n        ↓\n滞后与错配 / 断点\n        ↓\n斗争与反思\n        ↓\n对抗性适应\n        ↓\n双向靠拢与社会化\n        ↓\n冲突均衡（暂时、可再打开）")
    add_para(doc, "把这条链放回 ALIS，可以得到一个不把 AI 神秘化的解释：代码、模型、服务器、权限和教育场景构成物质与制度条件；简单/详细模式、旧类型码与新维度之间的漂移表现为滞后与错配；反馈、测试和作者的反复定标构成斗争与反思；人把 AI 纳入工作同时限制它，AI 又把人的概念变成可运行对象，这是一种对抗性适应和双向靠拢；当前可运行的基线只是冲突均衡，不是终极完成。")
    add_heading(doc, "7.2 三种需要避免的误读", 2)
    add_table(
        doc,
        ["误读", "为什么不符合 ALIS/《The Final》的对齐", "更稳妥的写法"],
        [
            ("技术决定论", "把模型能力当作研究方向的原因，忽略作者/场景/制度", "模型提供可能性，方向由条件与人的斗争共同形成。"),
            ("抽象人文主义", "把人说成天然主体，却不记录密钥、权限、代码和平台的物质约束", "主体性在具体技术和社会关系中被争取、被限制、被重新组织。"),
            ("流畅即真理", "把 Claude/DeepSeek 的完整语气当成证据", "把可读生成与事实、理论、现场经验分层。"),
        ],
        widths=[3.0, 7.0, 6.2],
        font_size=7,
    )
    add_heading(doc, "7.3 对作者主体性的保护", 2)
    add_bullet(doc, "不把跨学科跳跃抹成中性产品语言：社会学、民族志、HCI、软件工程和 AI 协作之间的跳跃本身就是思想材料。")
    add_bullet(doc, "不把未完成当作失败证据：简单/详细分叉、Interview OS 的延后、视觉和部署未再核验，都应作为历史条件写出来。")
    add_bullet(doc, "不让 AI 代替作者的最终判定：AI 可以提供结构、反例和实现，但“这个问题是否值得问”“这个解释是否负责任”仍属于作者与促进者。")
    add_callout(doc, "对齐结论", "ALIS 的最好版本不是把人变成更高效的 prompt 操作者，而是让技术痕迹、社会关系和人的反思彼此可见，并允许冲突在下一轮被重新打开。", "EAF3EA")


def section_eight(doc: Document) -> None:
    add_heading(doc, "八、当前技术地图、风险和作者审阅问题", 1)
    add_heading(doc, "8.1 现状地图", 2)
    add_table(
        doc,
        ["层", "已确认", "当前缺口", "建议审阅动作"],
        [
            ("内容", "2.2-draft；15 题、7 维；同时有 6/7 点量表", "版本命名和历史 README 漂移", "确定主量表并写 ADR，不直接改旧文本。"),
            ("评分", "纯函数、区间、偏差选择、XState", "简单/详细入口未统一", "先定产品与学术解释，再考虑映射。"),
            ("语言", "LiteLLM + DeepSeek 默认；坏输出 fallback", "真实服务、参数、成本与最新输出未复验", "以无个人数据 fixture 做受控 smoke。"),
            ("访谈", "六阶段 UI、草稿和保存 API", "浏览器视觉与长文本体验未由本次审计验收", "按阶段做真实浏览器 QA。"),
            ("数据", "SQLite/WAL、bundle、manifest/hash、owner 检查", "导出复制音频边界需重测", "用合成数据做隐私测试。"),
            ("托管", "根 `.gitignore`、MIT、README、审计稿已新增", "GitHub CLI token 失效，远程未确认", "重新认证后创建 private repo 并验证 SHA。"),
        ],
        widths=[2.0, 5.0, 5.0, 4.4],
        font_size=7,
    )
    add_heading(doc, "8.2 风险分层", 2)
    add_bullet(doc, "P0：历史 DeepSeek key 和部署信息在本地敏感档案中存在；Git 忽略只能阻止新提交，不能撤回已经泄露的凭据。建议项目所有者在控制台轮换。")
    add_bullet(doc, "P0：私有 GitHub 首次托管依赖新的认证；在 401 未解决前，不把上传说成已完成。")
    add_bullet(doc, "P1：简单/详细模式可能导致论文、用户界面和导出字段对不上；先决策，后代码。")
    add_bullet(doc, "P1：bundle 可能按指针复制音频；需要用合成数据重新核验最小化与删除。")
    add_bullet(doc, "P2：前端大 chunk、浏览器视觉、真实音频和公网部署属于未完成 QA，不应被 104/141 个单测覆盖替代。")
    add_heading(doc, "8.3 给作者的审阅问题", 2)
    for q in [
        "我是否同意把简单模式视为当前主要成品，把详细模式标作研究/工程内核，而不是把二者强行说成同一条产品？",
        "论文/报告最终要解释 7 点简单分类、6 点详细剖面，还是两层并存？两层并存时，哪个是主张、哪个是探索？",
        "六阶段访谈中的“反思”是测量结果、促进者方法，还是二者之间的转换场？",
        "哪些 Claude Code 交互是作者明确批准的，哪些只是模型建议？哪些失败应保留为方法材料？",
        "DeepSeek 的语言生成是否只需要可读性，还是需要结构化 JSON、版本化 prompt 和更强的审计？",
        "《The Final》的思想链中，ALIS 最有原创性的断点是量表、访谈、AI 协作本身，还是三者之间的冲突均衡？",
    ]:
        add_bullet(doc, q)


def appendix(doc: Document) -> None:
    add_heading(doc, "附录：证据索引与重建模板", 1)
    add_heading(doc, "A.1 主要证据", 2)
    add_table(
        doc,
        ["证据", "作用", "限制"],
        [
            ("`ALIS技术路径思辨总纲_v1.md`", "恢复研究目标、铁律、边界和早期路线", "是设计文本，不等于全部实现。"),
            ("S0-S6 交付报告", "恢复阶段意图和当时测试快照", "数字属于历史时点；“通过”范围需重新定义。"),
            ("`alis-mvp/frontend/src`", "确认模式、评分、状态机、离线与 UI 入口", "当前代码无法自动恢复作者当时的心理理由。"),
            ("`alis-mvp/backend/alis`", "确认 API、认证、LLM/fallback、数据库和导出", "真实外部服务状态不在本地代码中。"),
            ("723–727 反馈", "恢复外部测试、目标重述和代码治理经验", "反馈不是自动授权，也不等于已实现。"),
            ("原始交互 DOCX", "保留用户/Claude/工具事件的长程协作痕迹", "含敏感信息、压缩重复和工具 I/O 截断；不入 GitHub。"),
        ],
        widths=[5.0, 7.2, 5.0],
        font_size=7,
    )
    add_heading(doc, "A.2 后续每轮建议记录的最小字段", 2)
    add_code(doc, "turn_id / date / actor\nuser_goal / evidence_read / decision_boundary\nproposed_change / approved_scope / files_changed\ntests_run / result / unresolved\nmodel_role / prompt_version / fallback\nresearch_interpretation / fact_or_inference / author_review")
    add_para(doc, "如果继续使用 Claude Code 或 DeepSeek，建议保留上述字段的简短记录。它不要求保存所有隐藏思考，也不要求把个人数据复制到公开仓库；它只让人能回看一个决定如何从研究意图变成代码，又如何被测试和现场反馈重新改变。")
    add_heading(doc, "A.3 公开背景资料", 2)
    add_para(doc, "本文关于模型与工具的一般背景只依赖官方公开页面，访问时间为 2026-09-09；这些链接用于解释能力边界，不是 ALIS 项目日志。")
    for label, url in [
        ("DeepSeek-V3 官方仓库", "https://github.com/deepseek-ai/DeepSeek-V3"),
        ("DeepSeek Chat Completions API", "https://api-docs.deepseek.com/api/create-chat-completion/"),
        ("DeepSeek JSON Output", "https://api-docs.deepseek.com/guides/json_mode"),
        ("Anthropic prompt engineering guidance", "https://docs.anthropic.com/en/docs/build-with-claude/prompt-engineering/prompt-templates-and-variables"),
        ("Claude Code CLI usage", "https://docs.anthropic.com/en/docs/claude-code/cli-usage"),
        ("Claude Code memory and long-running context", "https://docs.anthropic.com/en/docs/claude-code/memory"),
    ]:
        add_bullet(doc, f"{label}：{url}")
    add_para(doc, "最后的事实边界：这份审阅稿的完成，证明的是“项目证据已经被重新组织并可审阅”，不是证明所有历史交互、线上部署或模型输出已经被完整恢复。", first_line=False, space_after=0)


def main() -> None:
    doc = Document()
    setup_styles(doc)
    for section in doc.sections:
        configure_section(section)

    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(header.add_run("ALIS · 人机民族志重建 / 审阅版"), color="6B7C89", size=8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(footer.add_run("2026-09-09  ·  仅供项目作者审阅  ·  第 "), color="6B7C89", size=8)
    add_page_number(footer)

    title_page(doc)
    add_contents(doc)
    doc.add_page_break()
    section_one(doc)
    doc.add_page_break()
    section_two(doc)
    doc.add_page_break()
    section_three(doc)
    doc.add_page_break()
    section_four(doc)
    doc.add_page_break()
    section_five(doc)
    doc.add_page_break()
    section_six(doc)
    doc.add_page_break()
    section_seven(doc)
    section_eight(doc)
    doc.add_page_break()
    appendix(doc)

    # Set core metadata without including secrets or absolute local paths.
    doc.core_properties.title = "ALIS AI 素养调查：人机民族志重建与审阅稿"
    doc.core_properties.subject = "项目迭代、prompt/loop、Claude Code、DeepSeek 与作者思想对齐"
    doc.core_properties.author = "ALIS contributors"
    doc.core_properties.comments = "Fact, historical record, inference and unknowns are explicitly separated."
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
