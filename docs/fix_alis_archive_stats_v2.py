from __future__ import annotations

from collections import Counter

from docx import Document

import build_alis_native_archive as builder
from build_alis_ethnography import set_cell_text


def main() -> None:
    source_doc = Document(builder.SOURCE)
    stats: Counter[str] = Counter()
    for paragraph in source_doc.paragraphs:
        builder.redact(paragraph.text, stats)

    output_doc = Document(builder.OUT)
    table = output_doc.tables[3]
    values = [
        stats.get("api_key", 0),
        stats.get("bearer_token", 0),
        stats.get("secret_assignment", 0),
        stats.get("private_key_block", 0),
        stats.get("ip_address", 0),
        stats.get("control_character", 0),
    ]
    for row, value in zip(table.rows[1:], values):
        set_cell_text(row.cells[1], str(value), size=8)
    output_doc.save(builder.OUT)
    print(builder.OUT)
    print(f"redactions={dict(stats)}")


if __name__ == "__main__":
    main()
