from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from build_alis_ethnography import (
    add_bullet,
    add_callout,
    add_code,
    add_heading,
    add_page_number,
    add_para,
    add_table,
    configure_section,
    setup_styles,
    style_run,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).with_name("ALIS人机民族志重建_九编对齐增强版.docx")


def add_kicker(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(label + "：")
    style_run(r1, bold=True, color="355C7D", size=9.5)
    r2 = p.add_run(text)
    style_run(r2, color="37474F", size=9.5)


def add_chapter_card(doc: Document, title: str, question: str, carrier: str, bridge: str, boundary: str, return_to_people: str) -> None:
    add_heading(doc, title, 2)
    add_kicker(doc, "核心问题", question)
    add_kicker(doc, "ALIS 的具体承载物", carrier)
    add_kicker(doc, "桥梁", bridge)
    add_kicker(doc, "自我反驳与边界", boundary)
    add_kicker(doc, "回到人和实践", return_to_people)


def title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(62)
    r = p.add_run("ALIS")
    style_run(r, bold=True, color="234E70", size=34)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(16)
    r = p2.add_run("人机民族志重建")
    style_run(r, bold=True, color="111111", size=21)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(20)
    r = p3.add_run("与《The Final》九编思想对齐审查和语言外围优化")
    style_run(r, bold=True, color="355C7D", size=14)
    add_callout(
        doc,
        "本稿性质",
        "这是 ALIS 既有审阅稿的派生增强版，不改动既有业务代码、成品或外部反馈原文。它把《The Final_优化版.docx》的九编结构作为思想校准器，逐编寻找 ALIS 的具体证据、必要桥梁、自我反驳和未决边界。",
        "EDF5F8",
    )
    add_callout(
        doc,
        "作者主体性",
        "对齐不是把 ALIS 写成《The Final》的附录，也不是用漂亮概念替换项目痕迹。凡是材料直接支持的内容写成事实，凡是把不同材料串联起来的内容写成分析，凡是无法确认的内容保留为未知。",
        "FFF5E6",
    )
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(70)
    r = p4.add_run("审查对象：The Final_优化版.docx（九编）\n工作对象：ALIS 人机民族志重建审阅稿\n版本：九编对齐增强版 / 2026-09-09")
    style_run(r, color="536878", size=10)
    doc.add_page_break()


def reading_note(doc: Document) -> None:
    add_heading(doc, "阅读说明与方法", 1)
    add_para(doc, "我先把《The Final_优化版.docx》作为一个完整的九编思想结构重读，再把它放回 ALIS 的代码、S0-S6 阶段报告、外部反馈和可见协作档案中。这里的“对齐”有一个很具体的含义：每一编都必须找到物质承载物，说明它如何通过社会关系和历史过程进入 ALIS，同时指出这一解释在哪些地方会失效。没有这三步，理论就只是贴在项目上的标签。")
    add_table(
        doc,
        ["证据层", "本稿用法", "语言纪律"],
        [
            ("项目事实", "代码、配置、测试、阶段报告和原始协作档案直接支持", "用“当前可见”“档案记录”“文件显示”"),
            ("The Final 概念", "九编中的问题意识、机制链和作者已形成的判断", "用作透镜，不把它冒充 ALIS 当时的原话"),
            ("分析性桥梁", "把工程事件、人的决定与概念连接起来", "使用“可以理解为”“更像”“可能”并给出承载物"),
            ("未知项", "缺失的源 JSONL、未重跑的外部模型、未再验证的部署", "不以流畅叙事填空，保留待作者核验"),
        ],
        widths=[2.5, 7.5, 6.2],
        font_size=7,
    )
    add_heading(doc, "全稿的对齐动作", 2)
    for item in [
        "保留 ALIS 自己的技术史：简单/详细模式、7 点/6 点量表、类型轴/连续维度、Interview OS 延后和当前测试基线都不被抹平。",
        "在每次理论跃迁前补一座桥：代码、内容包、数据库、权限、被试、促进者和反馈分别是什么，不让概念凭空升起。",
        "保留第一人称的决策位置，但不替作者编造缺失的心理动机；“我当时可能在做什么”必须标为重建，而非原话。",
        "让普通人的日常使用回到叙事：AI 素养不是熟练使用 prompt 的优越感，而是人在具体关系中能验证、拒绝、改写和承担后果。",
        "把“冲突均衡”写成暂时状态：能运行、能测试、能托管不等于理论完成；下一轮被试、甲方或作者仍可重新打开它。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "语言外围审查结论", 2)
    add_para(doc, "《The Final》的力量不在于每一段都抽象，而在于抽象之前有具体摩擦，抽象之后又回到关系、历史和普通人的生活。它允许长链条，但不允许跳步；允许自我修正，但不把未完成伪装成终极答案；允许干燥的幽默和突然的短句，但不以修辞代替证据。现有 ALIS 审阅稿的事实分层和工程地图是稳的，偏弱的地方是理论链条较集中在单一总纲，普通日常、作者第一人称和九编之间的逐章桥梁还不够细。本增强版只补这些桥，不把它改造成文学仿写。")
    add_code(doc, "具体经验 / 技术痕迹\n        ↓\n物质—社会—历史条件\n        ↓\n滞后、错配与断点\n        ↓\n斗争、反思与对抗性适应\n        ↓\n双向靠拢、社会化与权力迁移\n        ↓\n冲突均衡：暂时、可验证、可被下一轮重新打开")


def overview_matrix(doc: Document) -> None:
    add_heading(doc, "九编对齐总表", 1)
    add_para(doc, "九编不是一条把 ALIS 从“低级”推到“高级”的直线。它们更像九个相互咬合的观察面：同一个简单/详细分叉，既是第二编的错配，也是第三编的对抗性适应、第六编的权力定位和第九编的冲突均衡。下表先给出总览，随后逐编展开。")
    add_table(
        doc,
        ["《The Final》一编", "核心概念", "ALIS 具体锚点", "必须保留的边界"],
        [
            ("第一编", "物质性、社会历史性", "YAML、代码、SQLite、PWA、网络、研究者与被试", "AI 不是悬浮主体；代码也不是中性容器"),
            ("第二编", "矛盾、错配、滞后、断点", "简单/详细模式、旧 README、测试和反馈反复出现的分叉", "不是每个 bug 都是结构矛盾"),
            ("第三编", "斗争、反思、对抗性适应、双向靠拢", "S0-S6 gate、作者批准、Claude Code patch、定向测试", "“AI 协助”本身不能证明反思"),
            ("第四编", "物质、反思、通用、具身", "确定性评分、LLM 语言层、PWA/离线、真实访谈", "当前只有部分具身；不把思维链叫反思"),
            ("第五编", "AI 社会化、经验 Token 化、交互 Agent 化", "session/response/state、追问 Agent、DeepSeek 渲染", "Token 化是中介，不自动等于异化"),
            ("第六编", "制度、权力、溶解、矛盾转移", "认证、owner、CSRF、prompt 默认、导出、GitHub 治理", "分布式不等于平等；私有托管也有权力"),
            ("第七编", "主体性、强弱场域、日常生活", "六阶段访谈、拒答/不确定锚点、被试自述", "类型码不是身份或诊断"),
            ("第八编", "社会化通用智能", "ALIS 作为小型社会情境测试，不是 AGI 证明", "不作 AGI 能力宣传"),
            ("第九编", "冲突均衡、普惠、教育、自由", "MIT、私有 GitHub、审计、学生调查、人类复核", "托管基线不是普惠完成"),
        ],
        widths=[2.0, 4.2, 6.5, 4.1],
        font_size=6.8,
    )


def chapter_one(doc: Document) -> None:
    add_chapter_card(
        doc,
        "第一编　世界是什么：物质性、社会历史性",
        "一个 AI/人机系统存在之前，哪些物质、制度、关系和历史条件已经在场？",
        "ALIS 的内容 YAML、前端 React/TypeScript、后端 FastAPI/SQLAlchemy、SQLite/WAL、PWA 缓存、网络调用、认证 cookie、音频指针、研究者和被试。",
        "ALIS 的“智能”不是模型单独带来的属性。题目必须先被写成 YAML，分数必须被存成状态，追问必须经过 API、权限和 fallback，访谈必须在一个具体的学生调查场景中发生。模型输出只有进入这些关系，才成为研究事件。",
        "如果只把代码列出来，仍可能把物质性误读成技术清单；如果只谈人的意义，又会忘记 cookie、数据库、网络和密钥会直接限制谁能参与。当前项目也不能证明 PWA 离线、音频治理和线上部署在今天仍然有效。",
        "让被试看到的不只是一个类型码，而是一个可以追问、反驳和留下叙事的界面；让研究者知道每个结论经过了什么数据和权限路径。",
    )
    add_para(doc, "这一编对 ALIS 的第一个修正是：把“研究基础设施”说得更物质一些。量表不是纯思想，代码也不是思想的透明管道。一个 6 点或 7 点量表的选择，会进入校验、分数、类型、报告和被试对自己的理解。也因此，所谓人的主体性必须从这些具体载体中争取，而不是在系统外预先宣布。")


def chapter_two(doc: Document) -> None:
    add_chapter_card(
        doc,
        "第二编　为什么发展：矛盾、错配、滞后性与断点",
        "发展为何不是平滑升级，而是在不同速度、尺度和权力之间被错配逼出来？",
        "默认简单模式与详细模式并存；7 点简单量表与 6 点详细量表并存；旧 README、S0-S6 目标、反馈要求与当前入口存在时间差。",
        "当研究目标、产品入口和工程内核不同步时，补丁会暂时修复表面，却把差异推迟到结果页、追问、导出或论文解释中。反复出现的分叉，才可能从“bug”升级为结构性错配的证据；单次报错不足以承担这个判断。",
        "《The Final》里的“断点”不是给所有失败加意义。ALIS 当前只能证明存在版本漂移和入口分叉，不能凭此断定某一个历史决定必然由某种理论矛盾造成。需要把重复反馈、评价标准、责任归属和资源限制一起查回来。",
        "把简单/详细分叉写成可审阅的历史问题：哪个入口服务谁，哪种解释被暂时牺牲，哪种兼容性被保存，以及作者何时决定不再继续加补丁。",
    )
    add_para(doc, "这里可以保留一点作者式的自嘲：一个月后最难回忆的，往往不是某个函数怎么写，而是当时为什么容忍了两套刻度。代码把决定保存成事实，却不会自动保存决定当时的压力。人机民族志因此要记录滞后本身，而不是只在最终 README 里把它改写成一条整齐路线。")


def chapter_three(doc: Document) -> None:
    add_chapter_card(
        doc,
        "第三编　如何发展：斗争、反思、对抗性适应与双向靠拢",
        "人和 AI 如何在不抹平差异的情况下共同改变任务、流程和彼此的能力？",
        "S0-S6 的阶段 gate、作者对范围的批准、Claude Code 的只读盘点/小步修改/测试回报、反馈文件提出的“先重述目标再决定 patch 或重写”。",
        "真正的对抗性适应至少包含任务阻滞、限制 AI、改变工作流、改变证据和重新分配责任。ALIS 中，作者把 LLM 限定为语言层，把评分与状态机锁在确定性代码里；同时又接受 AI 的扫描、归纳和错误暴露来改变下一轮设计。",
        "档案中出现大量生成和工具调用，不等于每次都发生了反思。只有当输出改变了目标、证据、权限或责任，才有资格被写成一次反思性转折。Claude 的内部思考不能被直接当作作者的反思。",
        "让促进者和被试都能说“不对”：不接受一个流畅的追问，不接受一个类型码，不接受一个似乎已经完成的阶段报告，并把拒绝转成下一轮可追踪的修改。",
    )
    add_para(doc, "这一编是 ALIS 与《The Final》最直接的接合处。项目的技术铁律其实也是一种工作流斗争：把能生成语言的模型挡在评分、分支和研究编码之外。挡住它并不是拒绝技术，而是让双方在各自能承担的地方靠拢。作者提供方向和责任，AI 提供执行速度和反例；新的目标往往就在这次不完全重合中出现。")
    add_heading(doc, "可写入后续档案的最小事件单", 3)
    add_code(doc, "人提出目标 → AI 只读盘点 → AI 给方案 → 人批准范围\n→ AI 修改/测试 → 错误或反馈暴露新矛盾\n→ 人重新命名问题 → 接受、回滚、延后或重写")


def chapter_four(doc: Document) -> None:
    add_chapter_card(
        doc,
        "第四编　智能走向何处：物质性、反思性、通用性与具身性",
        "所谓智能，是否只是任务数量，还是必须能在现实反馈中改变自己与任务的关系？",
        "确定性评分/状态机与 LLM 语言层的分离；PWA、草稿、访谈保存、音频指针、浏览器界面和研究者现场。",
        "四种属性在 ALIS 里不是能力排行榜：物质性来自可运行的代码和存储，反思性来自断点后的改流程，通用性来自把同一研究目标带入技术/访谈/治理等不同社会情境，具身性来自回答、停顿、拒答、设备和现场反馈重新进入系统。",
        "当前 ALIS 有确定性评分和部分具身设计，但没有证据证明 AI 本身发生了反思，也没有资格据此宣称通用智能。不要把 Claude 的 chain-of-thought、DeepSeek 的长回答或多模态接口自动叫作反思性和具身性。",
        "把“会用 AI”改写成“能在现实反馈里重新判断”：被试可以修正自己的回答，促进者可以暂停追问，研究者可以修改目标，系统可以在异常时诚实 fallback。",
    )
    add_para(doc, "对齐后的 ALIS 不能只说“系统支持多线程、多模态、智能追问”。那样的词很容易越过证据。更稳妥的句子是：系统把某些现实反馈留在状态、文本和研究记录中；至于这些记录是否改变了人的判断、模型的行为和研究结论，需要单独设计研究。")


def chapter_five(doc: Document) -> None:
    add_chapter_card(
        doc,
        "第五编　人机如何彼此进入：AI 社会化、人的 Token 化与交互 Agent 化",
        "AI 如何进入社会关系，人的经验又如何被切分为可以被系统调用的状态？",
        "session、response、dimension score、profile、probe event、narrative segment、question reaction、bundle/export；六阶段访谈与后端追问 Agent；DeepSeek 作为运行时语言渲染器。",
        "一个人的回答先被编码为响应，再成为分数、画像、追问目标和叙事片段；这就是经验被 Token 化的技术中介。追问 Agent 则把题目、维度、骨架、状态和反馈接成持久的行动回路。AI 社会化不等于模型像小孩一样成长，而是它在拒绝、修复、反馈和制度约束中被允许进入关系。",
        "Token 化不是自动等于异化；它也可能让经验被保存、比较和重新谈论。反过来，记录也可能把无法进入 schema 的沉默、身体感和关系性经验排除在外。现有 ALIS 的代码能证明编码路径，不能单独证明被试感到被异化或被解放。",
        "给被试保留叙事和改写的入口，并在研究数据中标出“由人说出”“由规则计算”“由模型生成”“由促进者解释”的边界。",
    )
    add_para(doc, "这一编让 ALIS 的“访谈程序”比“智能问卷”更准确。系统不是把人一次性翻译成一个类型，而是让经验在回答、评分、追问、反应和叙事之间多次变形。每次变形都可能丢失东西，也可能获得新的可见性。研究的任务不是消灭这种中介，而是把它展示出来。")


def chapter_six(doc: Document) -> None:
    add_chapter_card(
        doc,
        "第六编　权力如何改变：显性制度、隐性制度、溶解与矛盾转移",
        "当制度被拆成权限、默认值、提示词、接口和习惯后，权力消失了吗？",
        "Argon2、HttpOnly/SameSite cookie、CSRF Origin/Referer、owner 检查、限流、角色、prompt 设计哲学、fallback、导出 manifest/hash、私有 GitHub 与 MIT 许可证。",
        "权力从“谁拥有一个大机构”迁移到“谁能改默认值、定义 schema、查看导出、批准部署、拒绝反馈”。当显性制度被拆成代码和工作流，功能可能溶解但不会消失；它会被内化为 prompt、状态机、权限和研究者习惯。",
        "私有 GitHub 提高了可恢复性，却不是权力平等；owner 检查保护被试，也把判断权集中在账户；MIT 允许再使用，但真实数据、密钥和部署仍然必须受限。不能把开源、私有托管或安全措施写成权力已经消失。",
        "把“谁可以看、改、导出、拒答、重跑、删除”作为研究设计的一部分；当个人与 AI 的小矛盾被转移到人与平台/机构之间时，要让转移本身可见。",
    )
    add_para(doc, "对 ALIS 来说，GitHub 私有托管不是一个与思想无关的运维动作。它保存了代码的时间线，也建立了新的访问边界。真正的开放不是把所有数据公开，而是让有权参与的人可以检查、复现、分叉和质疑；真正的保护也不是把一切关起来，而是说明哪些材料为什么不能公开。")


def chapter_seven(doc: Document) -> None:
    add_chapter_card(
        doc,
        "第七编　人在其中是什么：主体性、社会化、强弱场域与日常生活",
        "主体性不是离开技术，而是在关系和技术之中重新判断的能力；它如何在强弱对抗场域里形成？",
        "六阶段访谈、反应记录、开放叙事、拒答/不确定锚点、促进者说明、普通学生调查场景以及简单模式的可读结果页。",
        "强场域中，题目、反馈或现实后果迫使人重新判断；弱场域中，AI 过于顺滑地确认人，反而可能让人停止追问。ALIS 的访谈阶段如果只追求自然流畅，就会成为弱场域；如果允许被试修正、拒绝和回到生活材料，它才可能成为主体性训练。",
        "一个分类码不能代表主体，一个“不确定”选项也不能自动带来主体性。当前材料还不足以证明普通学生在真实使用中经历了强弱场域转换，需要观察、访谈和反应数据。",
        "把日常问题带回来：被试怎样在学习、工作、写作、游戏或家庭关系里验证 AI；他是否有权说“这不符合我的经验”；研究者是否会把这种拒绝当作有效资料。",
    )
    add_para(doc, "这里是语言优化最需要加重的一处。ALIS 不能把“AI 素养高”写成更快、更会、更自信地得到答案。更接近《The Final》的表述是：AI 素养包含在具体社会关系中重新判断目标、验证结果、修改问题、接受反驳和承担后果的能力。它不高高在上，甚至有一点麻烦；但正是这种麻烦阻止了流畅性吞掉经验。")


def chapter_eight(doc: Document) -> None:
    add_chapter_card(
        doc,
        "第八编　AGI 究竟是什么：从阿尔法进化到社会化通用智能",
        "通用智能是否应以任务数量衡量，还是以跨社会情境处理未知、否定和重新定义的能力衡量？",
        "ALIS 作为一个小型社会情境测试：它把量表、追问、访谈、研究者判断和导出放在同一现场，但不把自己冒充 AGI 研究或 AGI 证明。",
        "如果把 ALIS 放进这一编，它的价值在于提供一个很小却具体的检验场：系统是否能承认未知、请求更多信息、接受人的否定、保持 target 与语言的区别，并在反馈后重新组织问题。这里的通用性指跨角色和社会情境，而不是无限任务枚举。",
        "ALIS 没有也不应宣称测出了 AGI。DeepSeek 的追问质量、Claude Code 的开发能力和六阶段 UI 都不能推出社会化通用智能；最多只能成为未来研究如何定义“可质疑、可验证、可拒绝、可协商”的材料。",
        "把 AI 放回可被质疑的共同环境：它可以帮人理解，却必须允许人核验、拒绝、修改、重跑和追问“为什么”。",
    )
    add_para(doc, "这一编对 ALIS 的作用不是抬高项目，而是压低夸张。一个认真做的小系统，恰恰可以用来说明“大智能”不能只靠 benchmark 讲述。它必须面对具体人、具体机构、具体错误和具体后果。否则所谓通用，只是把实验室里的顺利迁移成了宣传。")


def chapter_nine(doc: Document) -> None:
    add_chapter_card(
        doc,
        "第九编　未来社会：冲突均衡、AI 普惠、教育与人的自由发展",
        "未来不是无冲突的技术乌托邦，而是人、AI、制度和教育在不断重开矛盾中形成的暂时均衡。",
        "MIT 许可证、私有 GitHub 基线、审计报告、可读民族志档案、学生调查、研究者复核、开源/Git/Fork/Issue 的知识生产方式。",
        "ALIS 的普惠不能只意味着“人人拿到一个 AI 工具”。它还意味着人有条件检查、拒绝、修改、创造和承担使用后果。教育不能只训练 prompt 技巧，而要训练材料性、社会性、历史性、验证、目标重写和对抗性适应。",
        "当前私有仓库只是恢复与安全的底线，不代表开放知识已经普惠；MIT 只覆盖适用的代码许可，不替代第三方依赖、模型条款、个人数据与机构伦理。项目也不能从一套调查程序直接推演未来社会。",
        "把最终判断交还给未来的人：本稿是可被作者、被试、甲方和下一轮开发推翻的暂时结论。开放意味着留下可检查的痕迹，自由意味着有现实资源和社会关系去改变这些痕迹。",
    )
    add_para(doc, "九编的收束最适合用在这里：冲突均衡不是和稀泥，而是让差异在某个阶段能够共同运行，同时保留重新冲突的可能。ALIS 当前的代码基线、审计材料和私有托管构成一种工程均衡；它还需要真实被试、普通日常、学术审阅和新的反馈来检验。它没有最终答案，也不需要假装有。")
    add_callout(doc, "九编综合判断", "ALIS 的思想主线可以被更清楚地说成：人在具体的物质—社会—历史条件中，借助 AI 形成可追问的研究关系；双方在错配和断点中改变流程与边界；新的可运行状态只是暂时均衡，必须回到普通人的生活、权力和自由发展中接受下一轮反驳。", "EAF3EA")


def language_review(doc: Document) -> None:
    add_heading(doc, "语言与思想的外围优化规则", 1)
    add_para(doc, "以下规则用于继续修订 ALIS 人机民族志正文，也可用于作者之后写项目报告、论文或《The Final》的相关延伸。它们不是要求全文机械模仿某种腔调，而是把已形成的思想底色守住。")
    add_table(
        doc,
        ["写作动作", "应当保留", "应当避免", "适用于 ALIS 的落点"],
        [
            ("开头", "从错误、反馈、卡顿、一个被试回答或一个权限决定切入", "先给宏大定义再寻找例子", "先写简单/详细分叉或一次追问 fallback"),
            ("概念跃迁", "技术痕迹 → 社会关系 → 历史条件 → 概念", "从一个漂亮词跳到另一个漂亮词", "每次写清 YAML、状态、角色、接口或反馈"),
            ("第一人称", "承认作者的犹豫、批准、拒绝和改变", "替作者虚构动机或把 AI 归纳当自述", "“我能确认的是……我只能推测的是……”"),
            ("自我反驳", "指出理论可能过度解释、指标可能吞掉经验", "为了显得深刻而无限否定", "类型码不是身份；日志不等于现场"),
            ("普通人", "把概念返回学习、工作、家庭和具体机构", "把被试只当数据行", "谁能验证、拒绝、改写和承担后果"),
            ("结论", "暂时、可复查、可被下一轮推翻", "“完成”“解决”“终极答案”", "当前基线是冲突均衡，不是项目终局"),
        ],
        widths=[2.4, 5.0, 5.0, 5.1],
        font_size=7,
    )
    add_heading(doc, "建议的段落内部顺序", 2)
    add_code(doc, "先写：一个具体事件/文件/界面/决定\n再写：它如何牵连人的关系、权限和历史\n然后写：The Final 的概念如何照亮它\n马上补：这个解释在哪里不够\n最后回到：下一步实践、普通人和可复查证据")
    add_heading(doc, "本次增强版对既有 ALIS 稿的具体优化", 2)
    for item in [
        "从“单一思想链”扩展为“九编逐章桥梁”，每章都含具体承载物、解释桥、自我反驳和回到人的落点。",
        "将“Claude Code/DeepSeek 的特点”与项目事实分开：Claude Code 是开发过程中的工具代理，DeepSeek 是运行时追问的语言层，二者都不承担作者的研究判断。",
        "将“人机民族志”从角色描述推进到权力和事件描述：谁发起改变、谁批准、谁能拒绝、什么被保存、什么被折叠。",
        "保留对私有 GitHub 的双重判断：它是恢复和安全的工程底线，同时也是新的权限制度；不能把托管动作写成开放已经完成。",
        "把普通人的日常、强弱场域、拒答和不确定性写进 AI 素养定义，避免它退化成 prompt 熟练度或模型崇拜。",
    ]:
        add_bullet(doc, item)


def review_questions(doc: Document) -> None:
    add_heading(doc, "作者最终审阅清单", 1)
    add_para(doc, "这是给作者的审阅入口，不是要求现在修改代码。每一个“是/否/待定”都可以作为下一轮人机协作的起点。")
    questions = [
        "九编中的哪一编最接近 ALIS 的原创断点：量表与访谈的关系、人与 AI 的协作，还是代码与研究判断的冲突？",
        "第一编的物质承载物是否足够具体，是否还需要补充真实设备、网络、服务器、音频和被试现场的证据？",
        "第二编的简单/详细分叉，是暂时的历史妥协、两个产品层，还是需要最终统一的技术债？",
        "第三编中，哪些改变确实由作者批准并改变了工作流，哪些只是 Claude Code 的顺滑建议？",
        "第四至第五编中，是否有把模型的生成、思维链或多模态能力过度写成反思、具身或 Agent 化的风险？",
        "第六编中，谁拥有查看、修改、导出、删除和重跑的权限；私有仓库与 MIT 如何同时服务保护和知识开放？",
        "第七编中，真实的普通学生日常是否出现，还是目前只有研究者对他们的抽象预设？",
        "第八编是否被严格写成问题意识，而不是借 ALIS 对 AGI 作能力背书？",
        "第九编的“自由发展”是否落到了现实条件、教育安排、知识访问和拒绝权，而不是抽象愿景？",
        "我是否愿意让未来的被试、同事、开发者或 AI 依据新证据推翻本稿的暂时判断？",
    ]
    for q in questions:
        add_bullet(doc, q)
    add_callout(doc, "收束", "如果一个理论不能让人更容易指出系统哪里错、谁受到影响、谁可以拒绝以及下一步怎样验证，它就还没有真正进入 ALIS。", "EAF3EA")
    add_heading(doc, "证据索引", 1)
    add_table(
        doc,
        ["材料", "本稿使用方式", "未越过的边界"],
        [
            ("The Final_优化版.docx", "九编结构、概念链、语言和思想审查", "不改原文件，不把其内容伪装成 ALIS 历史原话"),
            ("ALIS技术路径思辨总纲_v1.md", "研究目标、LLM 铁律、访谈方法和社会学方向", "设计意图不自动等于实现"),
            ("S0-S6 交付报告", "恢复阶段 gate、评分/状态/追问/采集/导出过程", "阶段测试数字属于历史快照"),
            ("当前前后端源码", "确认简单/详细入口、评分、API、认证和数据结构", "代码不能自动恢复作者当时的动机"),
            ("723 至 727 反馈", "恢复外部测试和代码治理经验", "只读参考，原文不改，不自动授权变更"),
            ("原生交互 DOCX", "确认用户/Claude/思考/工具事件与长程协作形态", "当前可见档案的工具 I/O 有转写截断；源 JSONL 未找到"),
        ],
        widths=[4.8, 7.0, 5.0],
        font_size=7,
    )
    add_para(doc, "本稿是对齐后的审阅材料，不是对 The Final 或 ALIS 的最终定稿。它完成了九编的外围审查与桥梁补足；下一轮若要进入正文逐句编辑，应由作者先确认哪些推断可以保留为第一人称，哪些必须改成分析语气。", first_line=False, space_after=0)


def main() -> None:
    doc = Document()
    setup_styles(doc)
    for section in doc.sections:
        configure_section(section)
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(header.add_run("ALIS · The Final 九编对齐 / 审阅版"), color="6B7C89", size=8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(footer.add_run("2026-09-09  ·  作者审阅材料  ·  第 "), color="6B7C89", size=8)
    add_page_number(footer)

    title_page(doc)
    reading_note(doc)
    doc.add_page_break()
    overview_matrix(doc)
    doc.add_page_break()
    chapter_one(doc)
    chapter_two(doc)
    doc.add_page_break()
    chapter_three(doc)
    chapter_four(doc)
    doc.add_page_break()
    chapter_five(doc)
    chapter_six(doc)
    doc.add_page_break()
    chapter_seven(doc)
    chapter_eight(doc)
    doc.add_page_break()
    chapter_nine(doc)
    language_review(doc)
    doc.add_page_break()
    review_questions(doc)

    doc.core_properties.title = "ALIS 人机民族志重建：与 The Final 九编思想对齐审查"
    doc.core_properties.subject = "ALIS 项目的人机民族志、九编桥梁、语言与思想外围优化"
    doc.core_properties.author = "ALIS contributors"
    doc.core_properties.comments = "Derived review document; source project code and The Final source document remain unchanged."
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
