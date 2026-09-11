from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from build_alis_ethnography import (
    add_page_number,
    add_table,
    configure_section,
    setup_styles,
    style_run,
)
from build_alis_native_archive import (
    DATE_RE,
    LABEL_RE,
    SOURCE,
    add_mixed_text,
    add_source_paragraph,
    redact,
    role_key,
    scan_redactions,
    sha256,
)


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "ALIS日期主题分卷"
MANIFEST = OUT_DIR / "ALIS日期主题分卷_manifest.json"


# The source DOCX contains 18 date-marked export batches. The names below are
# navigation labels, not claims that the source had 18 independent calendar
# days. The source order and every batch are retained.
BLOCK_DEFINITIONS = [
    ("V01", "2026-07-10", "项目发起、研究设计与路线图", "项目发起研究设计与路线图", "开场需求、社会学研究目标、量表方向、基础路线图与本地交付。"),
    ("V02", "2026-07-11", "基础测试、平台外延与量表问题", "基础测试平台外延与量表问题", "真人基础测试后的平台扩展、量表锚点和类 MBTI 结果机制问题。"),
    ("V03", "2026-07-10", "采集架构、DeepSeek 接入与自动化内测", "采集架构DeepSeek接入与自动化内测", "API、SQLite、录音/采集链路、S5/S6 及机器自动化内测安排。"),
    ("V04", "2026-07-11", "评分锚点、PWA 更新与首轮交付", "评分锚点PWA更新与首轮交付", "评分锚点取舍、基础设计落实、PWA 集成和向师门报备的首轮交付。"),
    ("V05", "2026-07-10", "采集架构、DeepSeek 接入与自动化内测（续接批次）", "采集架构DeepSeek接入与自动化内测_续接批次", "与前述 2026-07-10 采集/内测主题相邻的续接保留批次，不做去重。"),
    ("V06", "2026-07-11", "服务器选型、合规评估与部署规划", "服务器选型合规评估与部署规划", "轻量服务器、云平台、成本、合规和后续课程化/可复用部署规划。"),
    ("V07", "2026-07-10", "采集架构、DeepSeek 接入与自动化内测（续接批次）", "采集架构DeepSeek接入与自动化内测_续接批次02", "同一日期主题的再次保留批次；批次身份由源段范围区分。"),
    ("V08", "2026-07-11", "公网部署、安全组与灰度测试", "公网部署安全组与灰度测试", "服务器安全组、80/443 端口、部署授权与师门内部灰度测试。"),
    ("V09", "2026-07-10", "采集架构、DeepSeek 接入与自动化内测（续接批次）", "采集架构DeepSeek接入与自动化内测_续接批次03", "同一日期主题的再次保留批次；不把重复视为可删除噪声。"),
    ("V10", "2026-07-11", "上线检测、公网发布与部署推进", "上线检测公网发布与部署推进", "最终检测、服务器上线、访问链路和基础可落地产品的连续推进。"),
    ("V11", "2026-07-10", "采集架构、DeepSeek 接入与自动化内测（续接批次）", "采集架构DeepSeek接入与自动化内测_续接批次04", "同一日期主题的再次保留批次；保留上下文续接带来的重复结构。"),
    ("V12", "2026-07-11", "权限治理、账户系统与可复用性", "权限治理账户系统与可复用性", "SSH/安全组异常、账户管理、科研数据管理和课程化复用的权限治理。"),
    ("V13", "2026-07-10", "采集架构、DeepSeek 接入与自动化内测（续接批次）", "采集架构DeepSeek接入与自动化内测_续接批次05", "同一日期主题的再次保留批次；源档案的重复留存完整进入本卷。"),
    ("V14", "2026-07-11", "产品完善、师门账号与阶段交付", "产品完善师门账号与阶段交付", "产品审阅、首批账号准备、师门内部使用和阶段成果递送。"),
    ("V15", "2026-07-10", "采集架构、DeepSeek 接入与自动化内测（续接批次）", "采集架构DeepSeek接入与自动化内测_续接批次06", "同一日期主题的再次保留批次；不以内容相似度裁剪源记录。"),
    ("V16", "2026-07-11", "最终公网部署、权限故障与修复", "最终公网部署权限故障与修复", "公网部署、端口/SSH 权限故障、Nginx/账户处理和修复性推进。"),
    ("V17", "2026-07-10", "采集架构、DeepSeek 接入与自动化内测（续接批次）", "采集架构DeepSeek接入与自动化内测_续接批次07", "同一日期主题的再次保留批次；与其他批次并置供研究比较。"),
    ("V18", "2026-07-11", "部署收尾、成果汇报与原生档案归档", "部署收尾成果汇报与原生档案归档", "师门成果汇报、部署收尾、阶段停更和原生交互档案研究化归档。"),
]


def add_header_footer(doc: Document, marker: str) -> None:
    for section in doc.sections:
        configure_section(section)
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(header.add_run(f"ALIS · {marker} · 日期主题分卷"), color="6B7C89", size=8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(footer.add_run("源档案保序导航副本 · 第 "), color="6B7C89", size=8)
    add_page_number(footer)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    add_mixed_text(p, title, size=20, color="234E70", bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    add_mixed_text(p, subtitle, size=10.5, color="536878", bold=True)


def add_note(doc: Document, text: str, *, color: str = "27343B") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.14
    p.paragraph_format.space_after = Pt(5)
    add_mixed_text(p, text, size=9.2, color=color)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    style_run(p.add_run(text), bold=True, color="111111", size=14 if level == 1 else 11)


def build_event_maps(paragraphs):
    event_numbers: list[int] = []
    roles: list[str | None] = []
    labels: list[str | None] = []
    event_number = 0
    current_role: str | None = None
    for paragraph in paragraphs:
        raw = paragraph.text
        match = LABEL_RE.match(raw.strip())
        label = match.group(0) if match else None
        if label is not None:
            event_number += 1
            current_role = role_key(label)
        if DATE_RE.match(raw.strip()):
            current_role = None
        event_numbers.append(event_number)
        roles.append(current_role)
        labels.append(label)
    return event_numbers, roles, labels, event_number


def find_date_blocks(paragraphs):
    heads: list[tuple[int, str]] = []
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if DATE_RE.match(text):
            date_text = text.replace("──", "").strip()
            date = date_text.replace("年", "-").replace("月", "-").replace("日", "")
            heads.append((index, date))
    if len(heads) != len(BLOCK_DEFINITIONS):
        raise ValueError(f"Expected {len(BLOCK_DEFINITIONS)} date blocks, found {len(heads)}")
    blocks = []
    for position, (start, date) in enumerate(heads):
        end = heads[position + 1][0] if position + 1 < len(heads) else len(paragraphs)
        block_id, expected_date, topic, slug, rationale = BLOCK_DEFINITIONS[position]
        if date != expected_date:
            raise ValueError(f"{block_id} expected {expected_date}, found {date}")
        blocks.append(
            {
                "id": block_id,
                "date": date,
                "topic": topic,
                "slug": slug,
                "rationale": rationale,
                "start": start,
                "end": end,
            }
        )
    return blocks


def block_counts(paragraphs, start: int, end: int) -> tuple[Counter[str], Counter[str]]:
    labels: Counter[str] = Counter()
    roles: Counter[str] = Counter()
    for paragraph in paragraphs[start:end]:
        match = LABEL_RE.match(paragraph.text.strip())
        if match:
            label = match.group(0)
            labels[label] += 1
            roles[role_key(label) or label] += 1
    return labels, roles


def add_volume_front_matter(doc: Document, block, source_hash: str, paragraphs, event_numbers, labels) -> None:
    add_title(
        doc,
        f"ALIS 原生人机交互记录 · {block['date']}",
        f"{block['id']} · {block['topic']} · 脱敏研究导航副本",
    )
    add_note(
        doc,
        "本卷是从已确认的长档案生成的日期批次导航副本。它按源 DOCX 的原始段落顺序逐段收入，不去重、不合并相似段、不重写用户 Prompt 或 AI 输出；卷号、源段号和事件号是新增定位标记。原始长档案和原始协作 DOCX 均未被改写。",
        color="24445B",
    )
    add_note(
        doc,
        "日期是第一定位键；本卷的“原始批次”表示长档案中相邻日期线之间的保留片段，不等同于新的日历日。相同日期的重复/续接批次继续并置，这是交互导出、上下文续接和研究对象形成过程的一部分，不能仅凭文本相似度删除。",
        color="7A4B18",
    )
    label_counts, role_counts = block_counts(paragraphs, block["start"], block["end"])
    first_event = event_numbers[block["start"]]
    last_event = event_numbers[block["end"] - 1]
    redactions = Counter()
    for paragraph in paragraphs[block["start"] : block["end"]]:
        redact(paragraph.text, redactions)
    rows = [
        ("分卷编号", block["id"]),
        ("交互日期", block["date"]),
        ("主导主题", block["topic"]),
        ("原始批次", f"第 {int(block['id'][1:]):02d} 个日期线批次；保留源顺序"),
        ("源段范围", f"{block['start']:05d}–{block['end'] - 1:05d}（共 {block['end'] - block['start']:,} 段）"),
        ("全局事件号", f"{first_event:04d}–{last_event:04d}"),
        ("主题命名依据", block["rationale"]),
        ("源档案 SHA-256", source_hash),
        ("正文策略", "逐段保序；继承研究版脱敏；不删除、不去重、不声称恢复缺失 JSONL"),
        ("脱敏计数", ", ".join(f"{key}={value}" for key, value in sorted(redactions.items())) or "本卷未触发计数"),
    ]
    add_table(doc, ["字段", "值"], rows, widths=[3.1, 13.2], font_size=8)
    add_heading(doc, "本卷角色统计", 2)
    role_rows = [(key, f"{value:,}") for key, value in role_counts.most_common()]
    add_table(doc, ["归一角色", "事件标签数"], role_rows, widths=[7.5, 8.8], font_size=8)
    if label_counts:
        add_note(doc, "原始标签细分：" + "；".join(f"{key} {value:,}" for key, value in label_counts.most_common()), color="536878")
    add_heading(doc, "阅读提醒", 2)
    for item in [
        "用户消息、Claude 回复、思考、工具调用和工具结果均按现有档案提供的范围保留；工具 I/O 在源档案中若已截断，本分卷不补写缺失部分。",
        "中文正文使用 Microsoft YaHei；英文、命令、路径和 JSON 片段使用 Consolas；角色颜色只用于阅读导航，不表示事实等级或作者意图。",
        "正文中的脱敏标记表示源档案确有对应高风险值；本卷是研究副本，原始密钥、私钥、Bearer 凭据和 IP 不在研究副本中传播。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Pt(22)
        p.paragraph_format.first_line_indent = Pt(-10)
        p.paragraph_format.line_spacing = 1.1
        add_mixed_text(p, item, size=9.0, color="27343B")
    add_heading(doc, "正文：按源段保序收入", 1)
    add_note(doc, "以下正文仅增加“事件/源段”定位前缀和必要的日期/角色版式。源段文本本身保持原顺序；高风险值的可见替换沿用完整可读研究版的脱敏规则。", color="37474F")


def build_volume(block, paragraphs, source_hash, event_numbers, roles, labels) -> dict:
    output = OUT_DIR / f"{int(block['id'][1:]):02d}_{block['date']}_{block['slug']}.docx"
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {output}")
    doc = Document()
    setup_styles(doc)
    add_header_footer(doc, f"{block['id']} · {block['date']}")
    add_volume_front_matter(doc, block, source_hash, paragraphs, event_numbers, labels)
    local_redactions = Counter()
    for source_index in range(block["start"], block["end"]):
        raw = paragraphs[source_index].text
        safe = redact(raw, local_redactions)
        add_source_paragraph(
            doc,
            source_index,
            safe,
            roles[source_index],
            event_numbers[source_index],
            labels[source_index],
        )
    doc.core_properties.title = f"ALIS 原生人机交互记录 {block['date']} {block['topic']}"
    doc.core_properties.subject = "按日期和原始批次切分的保序人机交互研究副本"
    doc.core_properties.author = "ALIS contributors"
    doc.core_properties.comments = "Derived navigation copy. Source DOCX remains unchanged; no deduplication was performed."
    doc.save(output)
    return {
        "id": block["id"],
        "date": block["date"],
        "topic": block["topic"],
        "rationale": block["rationale"],
        "filename": output.name,
        "source_paragraph_start": block["start"],
        "source_paragraph_end_exclusive": block["end"],
        "source_paragraph_count": block["end"] - block["start"],
        "event_start": event_numbers[block["start"]],
        "event_end": event_numbers[block["end"] - 1],
        "label_counts": dict(block_counts(paragraphs, block["start"], block["end"])[0]),
        "redactions": dict(local_redactions),
    }


def build_index(blocks, volume_rows, paragraphs, source_hash, event_numbers, roles, labels, global_redactions) -> Path:
    output = OUT_DIR / "00_ALIS日期主题索引.docx"
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {output}")
    doc = Document()
    setup_styles(doc)
    add_header_footer(doc, "00 · 日期主题索引")
    add_title(doc, "ALIS 原生人机交互记录", "日期—主题分卷索引 · 脱敏研究导航版")
    add_note(
        doc,
        "本索引把 4,003 页长档案转成可按日期、原始批次和主导主题进入的多个阅读单元。分卷是新增研究副本，不回写、不替换、不删除原长档案；所有分卷共同覆盖源 DOCX 的 27,672 个段落。",
        color="24445B",
    )
    add_note(
        doc,
        "扫描显示，源档案实际出现两个交互日期（2026-07-10 与 2026-07-11），但内部有 18 条日期线。它们包含上下文续接、导出重复和阶段性保留结构，因此本轮采用“日期 + 原始批次 + 主导主题”的三重定位键，而不是把相似内容静默合并。",
        color="7A4B18",
    )
    add_heading(doc, "源档案与完整性边界", 1)
    global_rows = [
        ("源文件", SOURCE.name),
        ("源段落总数", f"{len(paragraphs):,}"),
        ("分卷覆盖", f"源段 00015–{len(paragraphs) - 1:05d} 进入 18 个日期批次；源段 00000–00014 收入本索引附录"),
        ("源档案 SHA-256", source_hash),
        ("日期识别", "2026-07-10、2026-07-11；18 个保留日期线批次"),
        ("处理方式", "保序切分；不去重、不合并、不修订正文；研究副本沿用高风险值脱敏"),
        ("证据边界", "当前目录未找到元数据所指向的完整 JSONL；工具 I/O 以现存 DOCX 为准"),
        ("GitHub 边界", "分卷正文保持本地；不会因本次生成自动进入 GitHub 远程仓库"),
    ]
    add_table(doc, ["字段", "值"], global_rows, widths=[3.1, 13.2], font_size=8)
    add_heading(doc, "日期—主题分卷清单", 1)
    table_rows = []
    for row in volume_rows:
        table_rows.append(
            (
                row["id"],
                row["date"],
                row["filename"],
                row["topic"],
                f"{row['source_paragraph_start']:05d}–{row['source_paragraph_end_exclusive'] - 1:05d}",
                f"{row['source_paragraph_count']:,}",
                f"{row['event_start']:04d}–{row['event_end']:04d}",
            )
        )
    add_table(
        doc,
        ["编号", "日期", "文件名", "主导主题", "源段范围", "段数", "事件范围"],
        table_rows,
        widths=[1.1, 2.0, 4.0, 4.0, 2.5, 1.2, 2.2],
        font_size=7,
    )
    add_heading(doc, "建议阅读路径", 1)
    for item in [
        "先读本索引，按日期筛选：先 2026-07-10 的研究设计/采集/内测线，再 2026-07-11 的测试、部署、权限和归档线。",
        "需要复原某一次具体交互时，以卷号、源段号和事件号共同定位；不要只引用主题名，因为主题名是本轮导航性编码。",
        "需要研究重复与续接时，平行比较同日的 V03/V05/V07/V09/V11/V13/V15/V17；这些卷是保留结构，不是内容错误的自动判定。",
        "需要研究人机民族志时，优先比较用户 Prompt、Claude 输出、思考、工具调用/结果和最终人类验收之间的落差，而不要把 AI 的中间文本直接等同于项目事实。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Pt(22)
        p.paragraph_format.first_line_indent = Pt(-10)
        p.paragraph_format.line_spacing = 1.1
        add_mixed_text(p, item, size=9.0, color="27343B")
    add_heading(doc, "全局脱敏统计", 1)
    redact_rows = [(key, str(value)) for key, value in sorted(global_redactions.items())]
    add_table(doc, ["类型", "替换次数"], redact_rows, widths=[7.5, 8.8], font_size=8)
    add_heading(doc, "源档案前置段附录", 1)
    add_note(doc, "为保证所有源段均可追溯，源段 00000–00014 以脱敏后的保序形式收入本索引；它们不是正文分卷的重复内容。", color="536878")
    preamble_redactions = Counter()
    for source_index in range(0, blocks[0]["start"]):
        raw = paragraphs[source_index].text
        safe = redact(raw, preamble_redactions)
        add_source_paragraph(doc, source_index, safe, roles[source_index], event_numbers[source_index], labels[source_index])
    doc.core_properties.title = "ALIS 原生人机交互记录 日期主题分卷索引"
    doc.core_properties.subject = "按日期、原始批次与主导主题定位长档案"
    doc.core_properties.author = "ALIS contributors"
    doc.core_properties.comments = "Navigation index for derivative date/topic volumes; source archive remains unchanged."
    doc.save(output)
    return output


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    source_hash = sha256(SOURCE)
    source_doc = Document(SOURCE)
    paragraphs = list(source_doc.paragraphs)
    event_numbers, roles, labels, total_events = build_event_maps(paragraphs)
    blocks = find_date_blocks(paragraphs)
    global_redactions = scan_redactions(paragraphs)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    volume_rows = []
    for block in blocks:
        row = build_volume(block, paragraphs, source_hash, event_numbers, roles, labels)
        volume_rows.append(row)
        print(f"created {row['filename']} paragraphs={row['source_paragraph_count']} events={row['event_end'] - row['event_start'] + 1}")
    index_path = build_index(
        blocks,
        volume_rows,
        paragraphs,
        source_hash,
        event_numbers,
        roles,
        labels,
        global_redactions,
    )
    manifest = {
        "source": str(SOURCE),
        "source_sha256": source_hash,
        "source_paragraph_count": len(paragraphs),
        "source_event_count": total_events,
        "source_redactions": dict(global_redactions),
        "coverage": {
            "index_preamble_start": 0,
            "index_preamble_end_exclusive": blocks[0]["start"],
            "volume_start": blocks[0]["start"],
            "volume_end_exclusive": len(paragraphs),
        },
        "volumes": volume_rows,
        "index_filename": index_path.name,
        "policy": "No source write-back, no deduplication, no content rewriting; derivative copies inherit canonical redaction.",
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"created {index_path.name}")
    print(f"manifest={MANIFEST}")
    print(f"source_sha256={source_hash}")
    print(f"source_paragraphs={len(paragraphs)} source_events={total_events}")
    print(f"source_redactions={dict(global_redactions)}")


if __name__ == "__main__":
    main()
